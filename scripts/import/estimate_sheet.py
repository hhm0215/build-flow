"""견적서 엑셀(.xlsx/.xls) 첫 시트 파서.

양식(0-based 행·열): 4행 C열=수신처, 5행 F열=현장명, 10행 G열=합계금액, 11행=머리글,
12행부터 품목(B=번호, C=품명, F/G=사이즈, H=규격→단위, I=수량, K=단가, M=금액, P=비고).
열 위치는 머리글 텍스트로 찾고, 찾지 못한 열만 고정 위치로 폴백한다.

서버(estimate-service)는 수량·단가를 소수 2자리까지 받고 항목 금액 = 수량×단가로 다시 계산한다.
그래서 여기서 계산한 합계(Σ 수량×단가)가 합계금액 셀과 다르면 ERROR로 둔다.

병합 셀: 머리글 위(수신처·현장명·합계 영역)는 병합 범위 전체에 왼쪽 위 값을 채운다. 품목 영역에서는
품명·규격·사이즈·비고처럼 글자 열만 세로 병합 아래 행에 값을 물려주고, 수량·단가·금액은 물려주지 않는다
(같은 금액을 두 번 세지 않기 위해). 물려받은 품명 행은 금액이 있으면 같은 품명의 별도 품목이 되고,
금액이 없으면 기존 규칙대로 "금액 없는 품명 행"으로 빠진다.
"""

from __future__ import annotations

import ast
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Optional

MAX_ROWS = 3000
SITE_MISSING_ERROR = "현장명 셀(R6C6)이 비어 있음"
NO_ITEMS_ERROR = "품목이 없음"
TOTAL_FORMULA_WARNING = "합계 셀 계산값 없음(수식) — 품목 합계로 검증 불가"
MAX_COLS = 40
MAX_QUANTITY = Decimal("99999999.99")
MAX_AMOUNT = Decimal("9999999999999.99")
LUMP_UNIT = "식"

# 고정 위치(폴백)
CLIENT_CELL = (4, 2)
SITE_CELL = (5, 5)
TOTAL_CELL = (10, 6)
YEAR_CELL = (3, 2)
HEADER_ROW = 11
FALLBACK_COLUMNS = {
    "no": 1, "name": 2, "size": 5, "spec": 7, "qty": 8,
    "total_qty": 9, "price": 10, "amount": 12, "note": 15,
}
HEADER_LABELS = {
    "NO": "no", "번호": "no", "품명": "name", "품목": "name", "사이즈": "size", "규격": "spec",
    "단위": "spec", "수량": "qty", "총수량": "total_qty", "단가": "price", "금액": "amount", "비고": "note",
}
SUMMARY_NAMES = {"합계", "소계", "총계", "계", "총합계", "합계금액", "부가세", "부가가치세", "VAT", "공급가액"}
CLIENT_SUFFIXES = ("귀하", "귀중", "님", "御中", "앞")
YEAR_PATTERN = re.compile(r"(?<!\d)((?:19|20)\d{2})\s*년")
DATE_PATTERN = re.compile(r"(20\d{2}|19\d{2})\s*[.\-/년]\s*(\d{1,2})\s*[.\-/월]\s*(\d{1,2})")
NUMBER_PATTERN = re.compile(r"-?\d[\d,]*(?:\.\d+)?")


class SheetFormatError(Exception):
    """양식을 읽을 수 없음(파일 형식·시트 없음)."""


@dataclass
class ParsedItem:
    row: int  # 0-based 시트 행
    name: str
    unit: str
    quantity: Decimal
    unit_price: Decimal

    @property
    def amount(self) -> Decimal:
        return self.quantity * self.unit_price


@dataclass
class ParsedEstimate:
    sheet: str
    client_name: str
    site_name: str
    cell_total: Optional[Decimal]
    estimate_date: Optional[dt.date]
    date_source: str
    sheet_year: Optional[int] = None  # 3행 근처 "YYYY년" 표기. 날짜로 쓰지 않는다
    total_formula_uncached: bool = False  # 합계 셀이 수식인데 저장된 계산값이 없음
    lump_from_total: bool = False  # 품목이 없어 합계금액 1식 품목으로 만듦
    items: list[ParsedItem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def computed_total(self) -> Decimal:
        return sum((item.amount for item in self.items), Decimal("0"))

    @property
    def row_span(self) -> str:
        if not self.items:
            return ""
        return f"{self.items[0].row + 1}-{self.items[-1].row + 1}"


# ─────────────────────────────────────────
# 셀 격자 읽기
# ─────────────────────────────────────────

@dataclass
class Sheet:
    title: str
    grid: list[list[Any]]
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)  # 0-based (r0, c0, r1, c1) 포함
    total_formula: Optional[Any] = None  # (r, c) -> 수식 여부 판별 함수(xlsx만)


def load_sheet(data: bytes) -> Sheet:
    """파일 바이트(매직 넘버로 형식 판별)에서 첫 시트의 값 격자와 병합 범위를 읽는다."""
    if data[:4] == b"PK\x03\x04":
        return _load_xlsx(data)
    if data[:8] == bytes.fromhex("D0CF11E0A1B11AE1"):
        return _load_xls(data)
    raise SheetFormatError("엑셀 파일 형식이 아닙니다(.xlsx/.xls 서명 없음)")


def load_grid(data: bytes) -> tuple[str, list[list[Any]]]:
    sheet = load_sheet(data)
    return sheet.title, sheet.grid


def _clip_merge(r0: int, c0: int, r1: int, c1: int) -> Optional[tuple[int, int, int, int]]:
    if r0 >= MAX_ROWS or c0 >= MAX_COLS:
        return None
    return r0, c0, min(r1, MAX_ROWS - 1), min(c1, MAX_COLS - 1)


def _load_xlsx(data: bytes) -> Sheet:
    import openpyxl

    # read_only 모드는 병합 범위를 주지 않으므로 일반 모드로 연다(파일 크기 상한은 호출 측에서 제한).
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    except Exception as exc:  # noqa: BLE001 - 손상 파일을 형식 오류로 통일
        raise SheetFormatError(f"xlsx를 열 수 없습니다({exc.__class__.__name__})") from None
    try:
        if not workbook.worksheets:
            raise SheetFormatError("시트가 없습니다")
        sheet = workbook.worksheets[0]
        max_row = min(sheet.max_row or 0, MAX_ROWS)
        grid = [list(row) for row in sheet.iter_rows(min_row=1, max_row=max_row, max_col=MAX_COLS,
                                                     values_only=True)] if max_row else []
        merges = []
        for rng in sheet.merged_cells.ranges:
            clipped = _clip_merge(rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1)
            if clipped:
                merges.append(clipped)
        return Sheet(sheet.title, grid, merges, lambda r, c: _xlsx_is_formula(data, r, c))
    finally:
        workbook.close()


def _xlsx_is_formula(data: bytes, r: int, c: int) -> bool:
    """data_only=True에서 비어 보이는 셀이 사실 계산값 없는 수식인지(수식 문자열로 다시 읽어 확인)."""
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception:  # noqa: BLE001
        return False
    try:
        sheet = workbook.worksheets[0]
        for row in sheet.iter_rows(min_row=r + 1, max_row=r + 1, min_col=c + 1, max_col=c + 1, values_only=True):
            value = row[0] if row else None
            if isinstance(value, str):
                return value.startswith("=")
            return type(value).__name__ in {"ArrayFormula", "DataTableFormula"}
        return False
    finally:
        workbook.close()


def _open_xls(data: bytes):
    import xlrd

    try:
        return xlrd.open_workbook(file_contents=data, on_demand=True, formatting_info=True)
    except Exception:  # noqa: BLE001 - 서식 정보를 못 읽는 파일은 병합 없이 다시 연다
        pass
    try:
        return xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as exc:  # noqa: BLE001
        raise SheetFormatError(f"xls를 열 수 없습니다({exc.__class__.__name__})") from None


def _load_xls(data: bytes) -> Sheet:
    import xlrd

    book = _open_xls(data)
    try:
        if book.nsheets < 1:
            raise SheetFormatError("시트가 없습니다")
        sheet = book.sheet_by_index(0)
        grid = []
        for r in range(min(sheet.nrows, MAX_ROWS)):
            row = []
            for c in range(min(sheet.ncols, MAX_COLS)):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    except Exception:  # noqa: BLE001
                        row.append(cell.value)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(None)
                else:
                    row.append(cell.value)
            grid.append(row)
        merges = []
        for rlo, rhi, clo, chi in getattr(sheet, "merged_cells", None) or []:
            clipped = _clip_merge(rlo, clo, rhi - 1, chi - 1)  # xlrd는 끝 행·열이 배타적
            if clipped:
                merges.append(clipped)
        return Sheet(sheet.name, grid, merges)
    finally:
        book.release_resources()


def apply_merges(grid: list[list[Any]], merges: list[tuple[int, int, int, int]], header_row: int,
                 text_cols: set[int]) -> set[tuple[int, int]]:
    """병합 값을 격자에 채운다. 품목 영역에서 물려받은 (행, 열)을 돌려준다."""
    inherited: set[tuple[int, int]] = set()
    for r0, c0, r1, c1 in merges:
        value = cell(grid, r0, c0)
        if value is None:
            continue
        for r in range(r0, min(r1, len(grid) - 1) + 1):
            row = grid[r]
            for c in range(c0, c1 + 1):
                if (r, c) == (r0, c0):
                    continue
                if r < header_row:
                    target = True
                else:
                    target = r0 > header_row and r > r0 and c in text_cols
                if not target:
                    continue
                if c >= len(row):
                    row.extend([None] * (c + 1 - len(row)))
                if row[c] is None:
                    row[c] = value
                    if r > header_row:
                        inherited.add((r, c))
    return inherited


def cell(grid: list[list[Any]], r: int, c: int) -> Any:
    if 0 <= r < len(grid) and 0 <= c < len(grid[r]):
        return grid[r][c]
    return None


# ─────────────────────────────────────────
# 값 변환
# ─────────────────────────────────────────

def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_label(value: Any) -> str:
    return re.sub(r"[\s.:·]", "", text(value)).upper()


class NumberError(ValueError):
    pass


def _plain(value: Decimal) -> Decimal:
    """뒤쪽 0을 없애되 지수 표기(1.5E+3)는 만들지 않는다."""
    value = value.normalize()
    return value.quantize(Decimal(1)) if value.as_tuple().exponent > 0 else value


def to_decimal(value: Any) -> Optional[Decimal]:
    """셀 값을 Decimal로. 빈 값·'-'는 None. 숫자로 읽을 수 없는 글자는 NumberError."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        # Excel 부동소수 잔차(예: 1234.5600000001)는 소수 6자리에서 정리한다.
        return _plain(Decimal(repr(round(value, 6))))
    raw = str(value).strip()
    if raw in {"", "-", "—"}:
        return None
    cleaned = re.sub(r"[\s,₩\\원]", "", raw)
    if re.fullmatch(r"-?\d+(\.\d+)?", cleaned):
        return _plain(Decimal(cleaned))
    raise NumberError(raw)


def quantity_with_unit(value: Any) -> Optional[Decimal]:
    """총수량 셀: 숫자이거나 '7.6M'·'4본'·'1EA'처럼 숫자 뒤에 단위가 붙은 글자. 그 밖의 글자는 None."""
    try:
        return to_decimal(value)
    except NumberError:
        pass
    match = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*[A-Za-z가-힣㎡㎥]{1,4}\s*", str(value))
    return _plain(Decimal(match.group(1))) if match else None


_EXPR_OPS = {ast.Add: lambda a, b: a + b, ast.Sub: lambda a, b: a - b,
             ast.Mult: lambda a, b: a * b, ast.Div: lambda a, b: a / b}


def _eval_expr(node: ast.AST) -> Decimal:
    if isinstance(node, ast.Expression):
        return _eval_expr(node.body)
    if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)) and not isinstance(node.value, bool):
        return Decimal(str(node.value))
    if isinstance(node, ast.BinOp) and type(node.op) in _EXPR_OPS:
        return _EXPR_OPS[type(node.op)](_eval_expr(node.left), _eval_expr(node.right))
    raise NumberError("식")


def quantity_from_text(value: Any) -> tuple[Optional[Decimal], Optional[str]]:
    """수량 셀: 숫자, '½', 또는 '(12.3*0.9)*2'처럼 사칙연산만 든 산출식 글자 → (값, 원래 식)."""
    try:
        return to_decimal(value), None
    except NumberError:
        pass
    raw = str(value).strip()
    if raw == "½":
        return Decimal("0.5"), None
    expr = raw.replace("×", "*").replace("x", "*").replace("X", "*").replace(",", "")
    if len(expr) > 100 or not re.fullmatch(r"[\d.\s+\-*/()]+", expr):
        raise NumberError(raw)
    try:
        result = _eval_expr(ast.parse(expr, mode="eval"))
    except (SyntaxError, InvalidOperation, ZeroDivisionError, NumberError):
        raise NumberError(raw) from None
    return _plain(result.quantize(Decimal("0.000001"))), raw


def amount_from_text(value: Any) -> Optional[Decimal]:
    """합계 셀: 숫자이거나 '일금 ... (₩1,234,000)' 같은 글자 안의 마지막 숫자 묶음."""
    try:
        return to_decimal(value)
    except NumberError:
        pass
    candidates = [m.group(0) for m in NUMBER_PATTERN.finditer(str(value))]
    if not candidates:
        return None
    try:
        return _plain(Decimal(candidates[-1].replace(",", "")))
    except InvalidOperation:
        return None


def fits_scale(value: Decimal, places: int = 2) -> bool:
    return value == value.quantize(Decimal(1).scaleb(-places)) if value.is_finite() else False


def strip_client_suffix(name: str) -> str:
    changed = True
    while changed and name:
        changed = False
        for suffix in CLIENT_SUFFIXES:
            if name.endswith(suffix) and len(name) > len(suffix):
                name = name[: -len(suffix)].rstrip(" :")
                changed = True
    return name.strip()


def find_date(grid: list[list[Any]], before_row: int) -> tuple[Optional[dt.date], str]:
    for r in range(min(before_row, len(grid))):
        for c, value in enumerate(grid[r]):
            if isinstance(value, dt.datetime):
                return value.date(), f"file-cell R{r + 1}C{c + 1}"
            if isinstance(value, dt.date):
                return value, f"file-cell R{r + 1}C{c + 1}"
            if isinstance(value, str):
                match = DATE_PATTERN.search(value)
                if match:
                    try:
                        return dt.date(*(int(g) for g in match.groups())), f"file-cell R{r + 1}C{c + 1}"
                    except ValueError:
                        continue
    return None, ""


def find_sheet_year(grid: list[list[Any]], header_row: int) -> Optional[int]:
    """3행 C열(없으면 2~4행)의 "YYYY년" 표기에서 연도만 뽑는다. 완전한 날짜 글자는 제외한다."""
    candidates = [cell(grid, *YEAR_CELL)]
    for r in range(2, min(5, header_row)):
        candidates.extend(grid[r] if r < len(grid) else [])
    for value in candidates:
        if not isinstance(value, str) or DATE_PATTERN.search(value):
            continue
        match = YEAR_PATTERN.search(value)
        if match:
            return int(match.group(1))
    return None


# ─────────────────────────────────────────
# 머리글·품목
# ─────────────────────────────────────────

def find_header(grid: list[list[Any]]) -> tuple[int, dict[str, int], bool]:
    """머리글 행과 열 매핑. 반환 (행, 매핑, 머리글 텍스트로 찾았는지)."""
    for r in range(min(len(grid), 40)):
        mapping: dict[str, int] = {}
        for c, value in enumerate(grid[r]):
            key = HEADER_LABELS.get(normalize_label(value))
            if key and key not in mapping:
                mapping[key] = c
        if "name" in mapping and ("amount" in mapping or "price" in mapping):
            return r, mapping, True
    return HEADER_ROW, dict(FALLBACK_COLUMNS), False


def _size_columns(mapping: dict[str, int]) -> list[int]:
    if "size" not in mapping:
        return []
    start = mapping["size"]
    later = sorted(c for c in mapping.values() if c > start)
    end = later[0] if later else start + 1
    return list(range(start, end))


def _merged_origin(merges: list[tuple[int, int, int, int]], r: int, c: int) -> tuple[int, int]:
    for r0, c0, r1, c1 in merges:
        if r0 <= r <= r1 and c0 <= c <= c1:
            return r0, c0
    return r, c


def parse_estimate(data: bytes, *, append_size: bool = False, append_note: bool = False,
                   lump_when_empty: bool = False, lump_name: str = "견적 합계") -> ParsedEstimate:
    sheet = load_sheet(data)
    sheet_name, grid = sheet.title, sheet.grid
    header_row, mapping, by_text = find_header(grid)
    warnings: list[str] = []
    errors: list[str] = []
    if not by_text:
        warnings.append("머리글을 찾지 못해 고정 열 위치를 사용함")
    else:
        missing = [k for k in ("qty", "price", "amount", "spec") if k not in mapping]
        for key in missing:
            mapping[key] = FALLBACK_COLUMNS[key]
        if missing:
            warnings.append("머리글에 없는 열은 고정 위치 사용: " + ", ".join(missing))

    size_cols = _size_columns(mapping)
    text_cols = {mapping["name"], mapping["spec"], *size_cols}
    if "note" in mapping:
        text_cols.add(mapping["note"])
    inherited = apply_merges(grid, sheet.merges, header_row, text_cols)
    inherited_name_rows = {r for r, c in inherited if c == mapping["name"]}

    client_name = strip_client_suffix(text(cell(grid, *CLIENT_CELL)))
    site_name = text(cell(grid, *SITE_CELL))
    raw_total = cell(grid, *TOTAL_CELL)
    cell_total = amount_from_text(raw_total) if raw_total is not None else None
    total_formula_uncached = False
    if cell_total is None:
        origin = _merged_origin(sheet.merges, *TOTAL_CELL)
        if raw_total is None and sheet.total_formula is not None and sheet.total_formula(*origin):
            total_formula_uncached = True
            warnings.append(TOTAL_FORMULA_WARNING)
        else:
            errors.append("합계금액 셀(R11C7)이 비었거나 숫자가 아님(수식은 엑셀에서 저장한 계산값이 필요)")

    estimate_date, date_source = find_date(grid, header_row)
    sheet_year = find_sheet_year(grid, header_row)

    items: list[ParsedItem] = []
    skipped_summary = 0
    skipped_blank_named = 0
    amount_rows: dict[int, Decimal] = {}  # 수량×단가≠금액 셀인 행 → 금액 셀 값
    for r in range(header_row + 1, len(grid)):
        name = text(cell(grid, r, mapping["name"]))
        try:
            qty, qty_expr = quantity_from_text(cell(grid, r, mapping["qty"]))
            total_qty = quantity_with_unit(cell(grid, r, mapping["total_qty"])) if "total_qty" in mapping else None
            price = to_decimal(cell(grid, r, mapping["price"]))
            amount = to_decimal(cell(grid, r, mapping["amount"]))
        except NumberError:
            errors.append(f"{r + 1}행: 수량·단가·금액에 숫자가 아닌 값")
            continue
        if not name:
            if amount not in (None, Decimal(0)) or price not in (None, Decimal(0)):
                errors.append(f"{r + 1}행: 품명 없이 금액/단가가 있음")
            continue  # 번호·0원만 있는 빈칸 행
        if normalize_label(name) in SUMMARY_NAMES:
            skipped_summary += 1
            continue
        if all(v in (None, Decimal(0)) for v in (qty, price, amount)):
            skipped_blank_named += 1
            continue
        if any(v is not None and v < 0 for v in (qty, price, amount)):
            errors.append(f"{r + 1}행: 음수 수량·단가·금액(할인 행 등)은 서버가 받지 않음")
            continue

        unit = text(cell(grid, r, mapping["spec"]))
        detail: list[str] = []
        if qty_expr:
            detail.append(f"산출 {qty_expr}")
            warnings.append(f"{r + 1}행: 수량 산출식 '{qty_expr[:40]}'을 계산해 {format_amount(qty)}로 넣음")
        if price not in (None, Decimal(0)) and qty not in (None, Decimal(0)):
            quantity, unit_price = qty, price
            if amount is not None and quantity * unit_price != amount:
                if total_qty and total_qty * unit_price == amount:
                    quantity = total_qty
                    warnings.append(f"{r + 1}행: 수량×단가≠금액, 총수량×단가=금액이라 총수량 사용")
                else:
                    warnings.append(f"{r + 1}행: 수량×단가({quantity * unit_price})≠금액 셀({amount})")
                    amount_rows[r] = amount
        elif amount not in (None, Decimal(0)):
            if unit and unit != LUMP_UNIT:
                detail.append(unit)  # 1식으로 바꾸므로 원래 규격은 품명에 남긴다
            quantity, unit_price, unit = Decimal(1), amount, LUMP_UNIT
        elif price is not None and qty not in (None, Decimal(0)):
            quantity, unit_price = qty, price  # 단가 0원(무상) 행
        else:
            quantity, unit_price, unit = Decimal(1), price or Decimal(0), LUMP_UNIT

        if (not fits_scale(quantity) or not fits_scale(quantity * unit_price)) and amount not in (None, Decimal(0))                 and fits_scale(amount):
            # 산출식·소수 수량이 소수 2자리를 넘으면 견적서 금액 셀을 1식으로 넣고 원래 수량은 품명에 남긴다
            detail.append(f"{format_amount(quantity)}{unit} × {format_amount(unit_price)}")
            warnings.append(f"{r + 1}행: 수량이 소수 2자리를 넘어 금액 셀 기준 1식으로 넣음")
            quantity, unit_price, unit = Decimal(1), amount, LUMP_UNIT
        if not unit:
            unit = LUMP_UNIT
        if len(unit) > 20:
            detail.append(unit)
            unit = LUMP_UNIT
            warnings.append(f"{r + 1}행: 규격이 20자를 넘어 단위 '식'으로 두고 품명에 덧붙임")
        if append_size:
            size = " ".join(t for t in (text(cell(grid, r, c)) for c in size_cols) if t)
            if size:
                detail.insert(0, size)
        if append_note and "note" in mapping:
            note = text(cell(grid, r, mapping["note"]))
            if note:
                detail.append(note)
        full_name = name + (f" ({', '.join(detail)})" if detail else "")
        if len(full_name) > 200:
            errors.append(f"{r + 1}행: 품명이 200자를 넘음")
            continue
        if not fits_scale(quantity) or not fits_scale(unit_price) or not fits_scale(quantity * unit_price):
            errors.append(f"{r + 1}행: 수량·단가·금액은 소수 2자리까지만 저장 가능")
            continue
        if quantity <= 0 or quantity > MAX_QUANTITY or unit_price > MAX_AMOUNT:
            errors.append(f"{r + 1}행: 수량·단가 범위 초과")
            continue
        items.append(ParsedItem(row=r, name=full_name, unit=unit, quantity=quantity, unit_price=unit_price))

    if amount_rows and cell_total is not None:
        by_amount = [item if item.row not in amount_rows else ParsedItem(
            row=item.row, name=f"{item.name} ({format_amount(item.quantity)}{item.unit} × {format_amount(item.unit_price)})"[:200],
            unit=LUMP_UNIT, quantity=Decimal(1), unit_price=amount_rows[item.row]) for item in items]
        if sum(i.amount for i in items) != cell_total and sum(i.amount for i in by_amount) == cell_total:
            # 견적서에 찍힌 금액 셀로 합계금액이 맞으면 그 행을 1식으로 바꾸고 원래 수량·단가는 품명에 남긴다
            items = by_amount
            warnings.append(f"수량×단가≠금액 행 {len(amount_rows)}개를 금액 셀 기준 1식으로 넣음(합계금액과 일치)")
    if skipped_summary:
        warnings.append(f"합계·소계 행 {skipped_summary}개 제외")
    if skipped_blank_named:
        warnings.append(f"금액 없는 품명 행 {skipped_blank_named}개 제외")
    inherited_items = sum(1 for item in items if item.row in inherited_name_rows)
    if inherited_items:
        warnings.append(f"병합된 품명을 아래 행 {inherited_items}개에 물려줌(각 행을 같은 품명의 별도 품목으로 넣음)")
    lump_from_total = False
    if not items and lump_when_empty and cell_total is not None and cell_total > 0:
        if fits_scale(cell_total) and cell_total <= MAX_AMOUNT:
            items.append(ParsedItem(row=TOTAL_CELL[0], name=(lump_name or "견적 합계")[:200], unit=LUMP_UNIT,
                                    quantity=Decimal(1), unit_price=cell_total))
            lump_from_total = True
            warnings.append("품목이 없어 합계금액을 1식 품목 하나로 넣음(--lump-when-empty)")
    if not items:
        errors.append(NO_ITEMS_ERROR)
    if not site_name:
        errors.append(SITE_MISSING_ERROR)
    if not client_name:
        warnings.append("수신처 셀(R5C3)이 비어 있어 거래처 없이 현장을 찾음")

    parsed = ParsedEstimate(
        sheet=sheet_name, client_name=client_name, site_name=site_name, cell_total=cell_total,
        estimate_date=estimate_date, date_source=date_source, sheet_year=sheet_year,
        total_formula_uncached=total_formula_uncached, lump_from_total=lump_from_total, items=items,
        warnings=warnings, errors=errors,
    )
    total = parsed.computed_total
    if total > MAX_AMOUNT:
        errors.append("품목 합계가 서버 한도를 넘음")
    if cell_total is not None and items and total != cell_total:
        errors.append(f"품목 합계({format_amount(total)}) ≠ 합계금액 셀({format_amount(cell_total)})")
    return parsed


def format_amount(value: Optional[Decimal]) -> str:
    if value is None:
        return "-"
    if value == value.to_integral_value():
        return f"{int(value):,}"
    return f"{value:,.2f}"
